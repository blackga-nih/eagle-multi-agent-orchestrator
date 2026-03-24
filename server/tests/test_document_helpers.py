"""Tests for document helper functions — extraction, normalization, and unfilled detection."""

import os
import sys
from unittest.mock import patch

import pytest

_server_dir = os.path.normpath(os.path.join(os.path.dirname(__file__), ".."))
if _server_dir not in sys.path:
    sys.path.insert(0, _server_dir)

ENV_PATCH = {
    "REQUIRE_AUTH": "false",
    "DEV_MODE": "true",
    "USE_BEDROCK": "false",
    "COGNITO_USER_POOL_ID": "us-east-1_test",
    "COGNITO_CLIENT_ID": "test-client",
    "EAGLE_SESSIONS_TABLE": "eagle",
    "USE_PERSISTENT_SESSIONS": "false",
    "S3_BUCKET": "test-bucket",
}


# ── Money Extraction ──────────────────────────────────────────────────

class TestExtractFirstMoneyValue:
    @pytest.fixture(autouse=True)
    def _import(self):
        with patch.dict(os.environ, ENV_PATCH, clear=False):
            from app.agentic_service import _extract_first_money_value
            self.fn = _extract_first_money_value

    def test_standard_currency(self):
        assert self.fn("Budget is $1,234,567.89 for FY26") == "$1,234,567.89"

    def test_simple_dollar_amount(self):
        assert self.fn("Cost: $500000") == "$500000"

    def test_shorthand_million(self):
        assert self.fn("We need about 1.2 million") == "1.2 million"

    def test_shorthand_k(self):
        assert self.fn("Roughly 500k budget") == "500k"

    def test_shorthand_m(self):
        assert self.fn("Estimate is 2.5M") == "2.5M"

    def test_dollar_sign_preferred_over_shorthand(self):
        result = self.fn("$750,000 which is about 750k")
        assert result == "$750,000"

    def test_empty_string(self):
        assert self.fn("") is None

    def test_none_input(self):
        assert self.fn(None) is None

    def test_no_money_value(self):
        assert self.fn("We need cloud hosting services") is None


# ── Period Extraction ──────────────────────────────────────────────────

class TestExtractPeriod:
    @pytest.fixture(autouse=True)
    def _import(self):
        with patch.dict(os.environ, ENV_PATCH, clear=False):
            from app.agentic_service import _extract_period
            self.fn = _extract_period

    def test_months(self):
        result = self.fn("Contract is 12 months base period")
        assert result is not None
        assert "12 months" in result

    def test_years(self):
        result = self.fn("Performance over 3 years")
        assert result is not None
        assert "3 years" in result

    def test_year_singular(self):
        result = self.fn("1 year base")
        assert result is not None
        assert "1 year" in result

    def test_trailing_context_captured(self):
        result = self.fn("12 months with two 6-month option periods")
        assert result is not None
        assert "12 months" in result
        assert len(result) <= 60  # Up to 40 chars trailing context

    def test_empty_string(self):
        assert self.fn("") is None

    def test_none_input(self):
        assert self.fn(None) is None

    def test_no_period(self):
        assert self.fn("We need cloud hosting") is None


# ── Section Bullets Extraction ─────────────────────────────────────────

class TestExtractSectionBullets:
    @pytest.fixture(autouse=True)
    def _import(self):
        with patch.dict(os.environ, ENV_PATCH, clear=False):
            from app.agentic_service import _extract_section_bullets
            self.fn = _extract_section_bullets

    def test_basic_heading_and_bullets(self):
        text = "Project Description:\n- Cloud hosting\n- DevOps support"
        result = self.fn(text)
        assert "project_description" in result
        assert result["project_description"] == ["Cloud hosting", "DevOps support"]

    def test_multiple_sections(self):
        text = (
            "Deliverables:\n- Monthly report\n- Final deliverable\n"
            "Security:\n- FedRAMP authorized\n- FISMA compliance"
        )
        result = self.fn(text)
        assert "deliverables" in result
        assert "security" in result
        assert len(result["deliverables"]) == 2
        assert len(result["security"]) == 2

    def test_star_bullets(self):
        text = "Technical Requirements:\n* Python 3.12\n* FastAPI"
        result = self.fn(text)
        assert "technical_requirements" in result
        assert "Python 3.12" in result["technical_requirements"]

    def test_bullets_without_heading_go_to_general(self):
        text = "- Standalone bullet\n- Another one"
        result = self.fn(text)
        assert "general" in result
        assert len(result["general"]) == 2

    def test_empty_input(self):
        assert self.fn("") == {}
        assert self.fn(None) == {}

    def test_quotes_stripped(self):
        text = 'Deliverables:\n- "Quoted item"'
        result = self.fn(text)
        assert result["deliverables"] == ["Quoted item"]


# ── Doc Type Normalization ─────────────────────────────────────────────

class TestNormalizeDocType:
    @pytest.fixture(autouse=True)
    def _import(self):
        with patch.dict(os.environ, ENV_PATCH, clear=False):
            from app.agentic_service import _normalize_create_document_doc_type
            self.fn = _normalize_create_document_doc_type

    def test_standard_type(self):
        assert self.fn("sow", "Statement of Work") == "sow"

    def test_ige_alias(self):
        assert self.fn("ige", "IGCE") == "igce"

    def test_statement_of_work_alias(self):
        assert self.fn("statement_of_work", "test") == "sow"

    def test_spaces_and_hyphens_normalized(self):
        assert self.fn("market-research", "Market Research") == "market_research"

    def test_sow_overridden_by_igce_title(self):
        assert self.fn("sow", "IGCE for Cloud Services") == "igce"

    def test_infer_from_title_when_empty(self):
        assert self.fn("", "Acquisition Plan for IT Services") == "acquisition_plan"

    def test_default_to_sow(self):
        assert self.fn("", "") == "sow"

    def test_none_input(self):
        assert self.fn(None, "") == "sow"

    def test_case_insensitive(self):
        assert self.fn("SOW", "test") == "sow"

    def test_acquisition_plan_type(self):
        assert self.fn("acquisition_plan", "Test AP") == "acquisition_plan"


# ── Infer Doc Type from Title ──────────────────────────────────────────

class TestInferDocTypeFromTitle:
    @pytest.fixture(autouse=True)
    def _import(self):
        with patch.dict(os.environ, ENV_PATCH, clear=False):
            from app.agentic_service import _infer_doc_type_from_title
            self.fn = _infer_doc_type_from_title

    def test_sow_title(self):
        assert self.fn("Statement of Work for Cloud Migration") == "sow"

    def test_igce_title(self):
        assert self.fn("Independent Government Cost Estimate") == "igce"

    def test_market_research_title(self):
        assert self.fn("Market Research Report for IT Services") == "market_research"

    def test_acquisition_plan_title(self):
        assert self.fn("Acquisition Plan for Network Upgrade") == "acquisition_plan"

    def test_justification_title(self):
        assert self.fn("Justification and Approval for Sole Source") == "justification"

    def test_empty_title(self):
        assert self.fn("") is None

    def test_none_title(self):
        assert self.fn(None) is None

    def test_unrecognized_title(self):
        assert self.fn("Random Document Title") is None


# ── Unfilled Template Detection ────────────────────────────────────────

class TestLooksLikeUnfilledTemplatePreview:
    @pytest.fixture(autouse=True)
    def _import(self):
        with patch.dict(os.environ, ENV_PATCH, clear=False):
            from app.agentic_service import _looks_like_unfilled_template_preview
            self.fn = _looks_like_unfilled_template_preview

    def test_empty_preview_returns_false(self):
        assert self.fn("sow", "") is False

    def test_placeholder_tokens_detected(self):
        assert self.fn("sow", "Title: {{PROJECT_TITLE}}\nScope: {{SCOPE_DESC}}") is True

    def test_sow_unfilled_markers(self):
        preview = (
            "This section should provide brief description of the project. "
            "The background information should identify the requirement in very general terms."
        )
        assert self.fn("sow", preview) is True

    def test_acquisition_plan_unfilled_markers(self):
        preview = (
            "This section should describe the requirement. "
            "Describe the competition strategy for this procurement. [TBD]"
        )
        assert self.fn("acquisition_plan", preview) is True

    def test_justification_unfilled_markers(self):
        preview = "[Contractor Name] will provide... [Provide detailed rationale for sole source]"
        assert self.fn("justification", preview) is True

    def test_mostly_filled_document_passes(self):
        # Schema-based: needs >=30% sections filled to pass.
        # SOW has 15 sections, so need at least 5 matching headings.
        preview = (
            "## 1. Background and Purpose\n"
            "NCI requires cloud hosting services for the EAGLE application.\n\n"
            "## 2. Scope\n"
            "The contractor shall provide AWS cloud infrastructure management.\n\n"
            "## 3. Period of Performance\n"
            "Base period of 12 months with two 6-month option periods.\n\n"
            "## 4. Place of Performance\n"
            "Work shall be performed at NCI facilities in Bethesda, MD.\n\n"
            "## 5. Tasks/Requirements\n"
            "Cloud migration, security hardening, CI/CD pipeline setup.\n\n"
            "## 6. Deliverables\n"
            "Monthly status reports, migration plan, final documentation.\n\n"
        )
        assert self.fn("sow", preview) is False

    def test_unknown_doc_type_without_schema_no_placeholders_passes(self):
        # For truly unknown types with no schema AND no {{PLACEHOLDER}} tokens,
        # the heuristic markers won't match either — should pass.
        # Patch out schema-based validation so we test the fallback path.
        with patch("app.template_registry.validate_document_completeness", return_value=None):
            assert self.fn("unknown_type_xyz", "Some normal content here") is False


# ── Field Name Normalization ───────────────────────────────────────────

class TestNormalizeFieldNames:
    @pytest.fixture(autouse=True)
    def _import(self):
        with patch.dict(os.environ, ENV_PATCH, clear=False):
            from app.template_registry import normalize_field_names
            self.fn = normalize_field_names

    def test_canonical_key_passes_through(self):
        result = self.fn({"description": "Test", "title": "SOW"}, "sow")
        assert result["description"] == "Test"

    def test_competition_type_aliased(self):
        result = self.fn({"competition_type": "Full and Open"}, "acquisition_plan")
        assert result.get("competition") == "Full and Open"
        assert "competition_type" not in result

    def test_estimated_cost_aliased_to_estimated_value(self):
        result = self.fn({"estimated_cost": "$500K"}, "acquisition_plan")
        assert result.get("estimated_value") == "$500K"

    def test_contractor_name_aliased(self):
        result = self.fn({"contractor_name": "Acme Corp"}, "justification")
        assert result.get("contractor") == "Acme Corp"

    def test_unknown_keys_preserved(self):
        result = self.fn({"custom_field": "value"}, "sow")
        assert result["custom_field"] == "value"

    def test_empty_data(self):
        result = self.fn({}, "sow")
        assert result == {}

    def test_unknown_doc_type_returns_unchanged(self):
        data = {"foo": "bar"}
        result = self.fn(data, "nonexistent_type")
        assert result == data

    def test_alias_does_not_overwrite_canonical(self):
        # If both canonical and alias are present, canonical wins
        result = self.fn(
            {"competition": "Full and Open", "competition_type": "Set-Aside"},
            "acquisition_plan",
        )
        assert result["competition"] == "Full and Open"

    def test_multiple_aliases_to_same_target(self):
        result = self.fn(
            {"vendor": "Acme", "vendor_name": "Acme Inc"},
            "justification",
        )
        # First alias wins via setdefault
        assert result.get("contractor") == "Acme"


# ── Markdown Sidecar Fallback ──────────────────────────────────────────

class TestMarkdownSidecarFallback:
    """Tests the api_get_document markdown sidecar preference over binary extraction."""

    @pytest.fixture(autouse=True)
    def _setup(self):
        import io
        from datetime import datetime
        from types import ModuleType
        from unittest.mock import MagicMock

        from botocore.exceptions import ClientError
        from fastapi import APIRouter
        from fastapi.testclient import TestClient

        self.io = io
        self.datetime = datetime
        self.MagicMock = MagicMock
        self.ClientError = ClientError
        self.APIRouter = APIRouter
        self.TestClient = TestClient

        with patch.dict(os.environ, ENV_PATCH, clear=False):
            fake_strands = ModuleType("app.strands_agentic_service")
            fake_streaming_routes = ModuleType("app.streaming_routes")

            async def _mock(*a, **kw):
                if False:
                    yield None

            fake_strands.sdk_query = _mock
            fake_strands.sdk_query_streaming = _mock
            fake_strands.MODEL = "test"
            fake_strands.EAGLE_TOOLS = []
            fake_streaming_routes.create_streaming_router = lambda *a, **kw: APIRouter()

            with patch.dict(
                sys.modules,
                {
                    "app.strands_agentic_service": fake_strands,
                    "app.streaming_routes": fake_streaming_routes,
                },
            ):
                import importlib

                for m in ("app.main", "app.changelog_store", "app.cognito_auth", "app.document_ai_edit_service", "app.spreadsheet_edit_service"):
                    sys.modules.pop(m, None)
                import app.main as main_module

                importlib.reload(main_module)

                import app.cognito_auth as _auth
                _auth.DEV_MODE = True

                self.app = main_module.app

                token = _auth.generate_test_token(user_id="dev-user", tenant_id="dev-tenant")
                self._auth_headers = {"Authorization": f"Bearer {token}"}

    def _build_s3_with_sidecar(self, sidecar_content: str | None):
        from docx import Document as DocxDocument

        s3 = self.MagicMock()
        doc = DocxDocument()
        doc.add_paragraph("Binary content")
        buf = self.io.BytesIO()
        doc.save(buf)
        docx_bytes = buf.getvalue()

        def get_object(*, Bucket, Key):
            if Key.endswith(".content.md"):
                if sidecar_content is not None:
                    return {
                        "Body": self.io.BytesIO(sidecar_content.encode()),
                        "ContentType": "text/markdown",
                        "ContentLength": len(sidecar_content),
                        "LastModified": self.datetime(2026, 3, 19, 1, 0, 0),
                    }
                raise self.ClientError(
                    {"Error": {"Code": "NoSuchKey", "Message": "Not found"}},
                    "GetObject",
                )
            if Key.endswith(".docx"):
                return {
                    "Body": self.io.BytesIO(docx_bytes),
                    "ContentType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    "ContentLength": len(docx_bytes),
                    "LastModified": self.datetime(2026, 3, 19, 1, 0, 0),
                }
            return {
                "Body": self.io.BytesIO(b"# Markdown"),
                "ContentType": "text/markdown",
                "ContentLength": 10,
                "LastModified": self.datetime(2026, 3, 19, 1, 0, 0),
            }

        s3.get_object.side_effect = get_object
        s3.generate_presigned_url.return_value = "https://signed.example/doc"
        return s3

    def test_sidecar_preferred_when_available(self):
        s3 = self._build_s3_with_sidecar("# SOW\n\nRich markdown content from AI")
        with patch("boto3.client", return_value=s3):
            with self.TestClient(self.app) as client:
                resp = client.get(
                    "/api/documents/eagle/dev-tenant/dev-user/documents/test.docx?content=true",
                    headers=self._auth_headers,
                )
        assert resp.status_code == 200
        data = resp.json()
        assert "Rich markdown content from AI" in data["content"]
        assert data["preview_mode"] == "markdown_sidecar"

    def test_falls_back_to_binary_when_no_sidecar(self):
        s3 = self._build_s3_with_sidecar(None)
        with patch("boto3.client", return_value=s3):
            with self.TestClient(self.app) as client:
                resp = client.get(
                    "/api/documents/eagle/dev-tenant/dev-user/documents/test.docx?content=true",
                    headers=self._auth_headers,
                )
        assert resp.status_code == 200
        data = resp.json()
        # Should have extracted binary content, not the sidecar
        assert data["preview_mode"] != "markdown_sidecar"
        assert data["is_binary"] is True
